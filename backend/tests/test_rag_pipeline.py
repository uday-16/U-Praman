"""Regression tests for grounding, document handling and the actual API contracts."""
import io
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
# Reuse the isolated authentication test collections; never touch real user data.
from test_login_sessions import database
from fastapi import FastAPI
from fastapi.testclient import TestClient
from app.routers import analysis, chat
from app.services import corpus, gemini_service
from app.services.document_reader import read_document
from app.schemas.analysis import RequirementInput
from app.services.ai_extractor import extract_requirements_from_input
from app.routers.auth import require_session

class RetrievalTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        docs = self.root / 'docs'; docs.mkdir()
        (docs / 'is.10500.1991.txt').write_text('Indian Standard\nDRINKING WATER SPECIFICATION\nFirst revision\nOld drinking water requirements for turbidity and potable water.', encoding='utf-8')
        (docs / 'is.10500.2012.txt').write_text('Indian Standard\nDRINKING WATER SPECIFICATION\nSecond revision\nDrinking water turbidity acceptable limit 1 NTU. Permissible limit 5 NTU.', encoding='utf-8')
        (docs / 'is.2925.1984.txt').write_text('Indian Standard\nINDUSTRIAL SAFETY HELMETS\nFirst revision\nSafety helmets provide impact protection for industrial workers.', encoding='utf-8')
        (docs / 'zIS10500Amd.1_2015.txt').write_text('AMENDMENT NO. 1 2015\nTO\nIS 10500 : 2012 DRINKING WATER\nRevise the drinking water testing procedure.', encoding='utf-8')
        (docs / 'is.15298.2.2011.txt').write_text('Indian Standard\nSAFETY FOOTWEAR\nFirst revision\nSafety footwear provides impact protection and penetration resistance.', encoding='utf-8')
        self.cache = patch.object(corpus, 'CACHE', self.root / 'cache'); self.cache.start()
        with patch.object(corpus, 'embedding_model', return_value=None): self.corpus = corpus.Corpus(docs)
    def tearDown(self): self.cache.stop(); self.temp.cleanup()
    def test_topical_and_no_match(self):
        hits = self.corpus.search('industrial helmets impact protection')
        self.assertEqual(hits[0]['number'], '2925')
        self.assertEqual(self.corpus.search('quantum banana spaceship'), [])
        self.assertEqual(self.corpus.search('IS 99999'), [])
    def test_product_ranking_excludes_different_safety_products(self):
        from app.schemas.analysis import ExtractedRequirement
        from app.services.vector_engine import rank_standards_for_requirement
        requirement=ExtractedRequirement(id='test',product_name='Industrial safety helmets',application='construction',purpose='impact protection',key_requirements=['impact protection and penetration resistance'],technical_parameters={},safety_parameters=[],extracted_at='2026-09-24')
        with patch('app.services.vector_engine.get_corpus',return_value=self.corpus):
            results=rank_standards_for_requirement(requirement)
        self.assertTrue(results)
        self.assertEqual({r.id for r in results},{'is-2925-1984'})

    def test_editions_and_amendments(self):
        hits = self.corpus.search('IS 10500 turbidity', 20)
        self.assertTrue(hits)
        self.assertNotIn('1991', [c['year'] for c in hits])
        amendment = next(d for d in self.corpus.documents if d['amendment'])
        self.assertEqual(amendment['base_year'], '2012')
        self.assertEqual({c['year'] for c in self.corpus.search('IS 10500:1991')}, {'1991'})
    def test_source_page_and_filter(self):
        hit = self.corpus.search('water', standard_id='is-10500-2012')[0]
        self.assertEqual(hit['page'], 1)
        self.assertIn('1 NTU', hit['text'])
        self.assertEqual(self.corpus.search('water', standard_id='unknown'), [])
    def test_index_changes_with_source(self):
        old = self.corpus.fingerprint
        (self.corpus.directory / 'is.10500.2012.txt').write_text('Drinking water changed source document with new testing requirements.', encoding='utf-8')
        with patch.object(corpus, 'embedding_model', return_value=None): fresh = corpus.Corpus(self.corpus.directory)
        self.assertNotEqual(old, fresh.fingerprint)

class DocumentTests(unittest.TestCase):
    def test_invalid_uploads(self):
        for data, name in [(b'', 'a.txt'), (b'not a pdf', 'a.pdf'), (b'word content', 'a.doc'), (b'\xff\xfe', 'a.txt')]:
            with self.subTest(name=name), self.assertRaises(ValueError): read_document(data, name)
    def test_docx_table_and_text(self):
        from docx import Document
        doc = Document(); doc.add_paragraph('Procure drinking water for municipal supply.')
        table = doc.add_table(rows=1, cols=2); table.cell(0,0).text='Turbidity'; table.cell(0,1).text='1 NTU'
        stream = io.BytesIO(); doc.save(stream)
        content = read_document(stream.getvalue(), 'tender.docx')[0]['text']
        self.assertIn('Turbidity | 1 NTU', content)
    def test_fallback_does_not_invent(self):
        with patch('app.services.ai_extractor.extract_requirements_with_gemini', side_effect=gemini_service.GenerationUnavailable('offline')):
            result = extract_requirements_from_input(RequirementInput(text='Supply industrial safety helmets.'))
        self.assertEqual(result.application, '')
        self.assertEqual(result.technical_parameters, {})
        self.assertNotIn('QCO', str(result.model_dump()))
        self.assertEqual(result.extraction_mode, 'source-text')

class GroundingTests(unittest.TestCase):
    def test_no_evidence_does_not_call_model(self):
        with patch.object(gemini_service, 'generate') as call:
            answer, mode = gemini_service.grounded_reply('invent an IS standard', [])
        call.assert_not_called(); self.assertEqual(mode, 'no-evidence')
    def test_invalid_citation_falls_back_to_source(self):
        citations = [{'is_number':'IS 10500','source':'water.pdf','page':4,'text':'Turbidity acceptable limit 1 NTU.'}]
        with patch.object(gemini_service, 'generate', return_value='Confirmed compliant [99]'):
            answer, mode = gemini_service.grounded_reply('water', citations)
        self.assertEqual(mode, 'extractive'); self.assertIn('1 NTU', answer); self.assertNotIn('Confirmed compliant', answer)

class AnalysisApiTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.store=patch.object(analysis,'STORE',Path(self.temp.name)); self.store.start()
        from app.routers import reports
        app=FastAPI(); app.include_router(analysis.router); app.include_router(chat.router); app.include_router(reports.router)
        self.app=app; self.client=TestClient(app)

    def tearDown(self): self.store.stop(); self.temp.cleanup()
    def authorize(self, user='test-officer'):
        self.app.dependency_overrides[require_session]=lambda: {'id':user}
    def test_authentication_required(self):
        self.assertEqual(self.client.post('/analysis/extract',json={'text':'Supply safety helmets'}).status_code,401)
    def test_missing_record_is_not_fabricated(self):
        self.authorize()
        self.assertEqual(self.client.get('/analysis/anl-missing').status_code,404)
        response=self.client.post('/analysis/confirm?extraction_id=req-missing',json={'product_name':'helmet','application':'','purpose':'','key_requirements':[]})
        self.assertEqual(response.status_code,404)
    def test_persistence_and_owner_isolation(self):
        self.authorize()
        with patch('app.services.ai_extractor.extract_requirements_with_gemini', side_effect=gemini_service.GenerationUnavailable('offline')):
            result=self.client.post('/analysis/extract',json={'text':'Supply industrial safety helmets.'})
        self.assertEqual(result.status_code,200)
        path='/analysis/extractions/'+result.json()['id']
        self.assertEqual(self.client.get(path).status_code,200)
        self.authorize('other-officer')
        self.assertEqual(self.client.get(path).status_code,404)
    def test_bad_upload(self):
        self.authorize()
        response=self.client.post('/analysis/upload',files={'file':('bad.pdf',b'not pdf','application/pdf')})
        self.assertEqual(response.status_code,422)

    def test_analysis_job_lifecycle(self):
        import time
        self.authorize()
        from app.services.job_service import create_analysis_job, load_job
        from app.schemas.analysis import RequirementInput
        # Create persistent job
        job = create_analysis_job(
            user_id='test-officer',
            input_type='text',
            requirement_title='Industrial Safety Helmets',
            raw_input=RequirementInput(text='Supply industrial safety helmets with chin strap and shock absorption.')
        )
        self.assertTrue(job.id.startswith('job-'))
        self.assertEqual(job.user_id, 'test-officer')
        self.assertEqual(job.total_stages, 8)

        # Query job via API
        resp = self.client.get(f'/analysis/jobs/{job.id}')
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(data['id'], job.id)
        self.assertIn('stages', data)
        self.assertEqual(len(data['stages']), 8)

        # Wait briefly for background execution to complete or advance
        for _ in range(30):
            updated = load_job(job.id, 'test-officer')
            if updated and updated.status in ('COMPLETED', 'FAILED'):
                break
            time.sleep(0.1)

        final_job = load_job(job.id, 'test-officer')
        self.assertIsNotNone(final_job)
        if final_job.status == 'COMPLETED':
            self.assertEqual(final_job.progress_percent, 100)
            self.assertIsNotNone(final_job.result)
            self.assertTrue(final_job.result.id.startswith('anl-'))
            self.assertGreater(len(final_job.completed_stages), 0)

        # Check owner isolation
        self.authorize('other-officer')
        self.assertEqual(self.client.get(f'/analysis/jobs/{job.id}').status_code, 404)

    def test_report_and_pdf_generation(self):
        self.authorize()
        with patch('app.services.ai_extractor.extract_requirements_with_gemini', side_effect=gemini_service.GenerationUnavailable('offline')):
            extract_res = self.client.post('/analysis/extract', json={'text': 'Supply industrial safety helmets with shock absorption.'})
        self.assertEqual(extract_res.status_code, 200)
        ext_id = extract_res.json()['id']

        confirm_res = self.client.post(f'/analysis/confirm?extraction_id={ext_id}', json={
            'product_name': 'Industrial safety helmets',
            'application': 'Construction',
            'purpose': 'Head protection',
            'key_requirements': ['Shock absorption', 'Penetration resistance']
        })
        self.assertEqual(confirm_res.status_code, 200)
        anl_id = confirm_res.json()['id']

        # Create report
        rep_res = self.client.post('/reports', json={'analysis_id': anl_id, 'officer_name': 'Test Officer'})
        self.assertEqual(rep_res.status_code, 200)
        report_data = rep_res.json()
        report_id = report_data['id']
        self.assertEqual(report_data['officer_name'], 'Test Officer')

        # Download PDF
        pdf_res = self.client.get(f'/reports/{report_id}/download')
        self.assertEqual(pdf_res.status_code, 200)
        self.assertEqual(pdf_res.headers['content-type'], 'application/pdf')
        self.assertTrue(pdf_res.content.startswith(b'%PDF'))
        self.assertGreater(len(pdf_res.content), 2000)

if __name__ == '__main__': unittest.main()


